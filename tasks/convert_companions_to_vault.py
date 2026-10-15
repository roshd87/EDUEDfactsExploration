#!/usr/bin/env python3
"""Convert the EDFacts companion documents into Obsidian notes under wiki/edfacts/.

- BRSI (Business Rules Single Inventory, SY 2025-26 v1.04 XLSX, the newest published):
  `BRSI/BRSI SY 2025-26.md` (scope, version history, cross-file rules) plus
  `<FSnnn>/<FSnnn> Business Rules.md` for every spec that has a note under wiki/edfacts/.
- EDFacts Workbook with FAQs (SY 2025-26 PDF): `Workbook/EDFacts Workbook SY 2025-26.md`.
- EDPass metadata user guides (PDF): the newest guide per collection from the metadata page,
  `Metadata/<Collection> Metadata User Guide.md`.

File-spec mentions (FS052, FS5017 ...) become wikilinks to the spec notes / metadata notes.
Run after `convert_specs_to_vault.py`, then run that script again so the spec notes pick up the
back-links (it reads the companion notes' frontmatter).

    /home/<you>/.venvs/edfacts-tools/bin/python tasks/convert_companions_to_vault.py

Needs openpyxl and pymupdf4llm (not in the system Python: `python3 -m venv ~/.venvs/edfacts-tools &&
~/.venvs/edfacts-tools/bin/pip install openpyxl pymupdf4llm`).
"""
from __future__ import annotations

import argparse
import collections
import csv
import html
import re
import sys
from pathlib import Path

import openpyxl
import pymupdf4llm

sys.path.insert(0, str(Path(__file__).resolve().parent))
from convert_specs_to_vault import (  # noqa: E402
    FILTER_YEAR, FROZEN, REF, ROOT, TARGET_YEAR, fix_bold, parse_title, read_inventory, read_manifest, spec_key,
)

OUT = ROOT / "wiki" / "edfacts"
BRSI_XLSX = "business-rules/edfacts-business-rules-single-inventory-sy-2025-26-v104-113313.xlsx"
BRSI_LABEL = "SY 2025-26 v1.04"
WORKBOOK_PDF = "workbook/edfacts-workbook-faqs-sy-2025-26-113626.pdf"
WORKBOOK_NOTE = "EDFacts Workbook SY 2025-26"
INDEX_NOTE = "EDFacts Specs"


# ----------------------------------------------------------------------------- shared


def spec_notes() -> dict[str, str]:
    """FS id -> spec note basename, for the specs converted by convert_specs_to_vault.py."""
    old = {spec_key(r) for r in read_inventory(FILTER_YEAR) if r["status"] in ("Active", "Optional")}
    out = {}
    for r in read_inventory(TARGET_YEAR):
        key = spec_key(r)
        if key in old and not key[1]:  # FS029 links to the Directory note, not the LA one
            name, _, _ = parse_title(r["title"])
            out[key[0]] = f"{key[0]} {name}"
    return out


def q(s: str) -> str:
    return '"' + str(s).replace('"', '\\"') + '"'


def frontmatter(fields: dict) -> str:
    lines = ["---"]
    for k, v in fields.items():
        if isinstance(v, list):
            lines.append(f"{k}: [" + ", ".join(v) + "]")
        elif isinstance(v, str) and (":" in v or v == "" or v[0] in "[{\"'*&!%@`#" or v != v.strip()):
            lines.append(f"{k}: {q(v)}")
        else:
            lines.append(f"{k}: {v}")
    lines.append("---")
    return "\n".join(lines)


class Linker:
    def __init__(self, specs: dict[str, str], meta: dict[str, str]):
        self.specs, self.meta = specs, meta

    def fs(self, fs_id: str, label: str | None = None, self_note: str = "", in_table: bool = False) -> str:
        label = label or fs_id
        target = self.specs.get(fs_id) or self.meta.get(fs_id)
        if not target or target == self_note:
            return label
        sep = "\\|" if in_table else "|"  # Obsidian needs the alias pipe escaped inside table cells
        return f"[[{target}{sep}{label}]]"

    def linkify(self, text: str, self_note: str = "", in_table: bool = False) -> str:
        """Turn FSnnn / FS5nnn mentions into wikilinks, skipping existing links and code spans."""
        out_lines = []
        for line in text.split("\n"):
            table = in_table or line.startswith("|")
            parts = re.split(r"(\[\[[^\]]*\]\]|`[^`]*`)", line)
            for i, p in enumerate(parts):
                if i % 2:
                    continue
                p = re.sub(r"\bFS\s?(\d{3,4})\b", lambda m: self.fs("FS" + m.group(1), m.group(0), self_note, table), p)
                if table:  # '|029|' in a table whose first column is the file spec
                    p = re.sub(r"^\|(\d{3})\|", lambda m: "|" + self.fs("FS" + m.group(1), m.group(1), self_note, True) + "|", p)
                parts[i] = p
            out_lines.append("".join(parts))
        return "\n".join(out_lines)


def source_block(rel: str, what: str) -> str:
    return (
        f"> [!info] Source\n> Converted by `tasks/convert_companions_to_vault.py` from the frozen SY {TARGET_YEAR} "
        f"catalog copy of {what} (`wiki/reference/edfacts/{TARGET_YEAR}/{rel}`). The original remains the authority."
    )


# ----------------------------------------------------------------------------- PDF -> markdown

HEADING = re.compile(r"^(#+)\s*(.*?)\s*$")
NUMBERED = re.compile(r"^(?:(\d+(?:\.\d+)*)\.?|Appendix [A-Z]:?)\s", re.I)


def pdf_markdown(pdf: Path) -> tuple[str, str]:
    """Return (front, body): body starts at the first numbered section heading."""
    chunks = pymupdf4llm.to_markdown(str(pdf), page_chunks=True)
    pages = [c["text"] for c in chunks]
    # Running headers/footers: short lines that repeat on many pages.
    freq = collections.Counter()
    for p in pages:
        freq.update({ln.strip() for ln in p.splitlines() if 0 < len(ln.strip()) < 90})
    boiler = {ln for ln, n in freq.items() if n >= max(3, len(pages) * 0.3)}
    lines = []
    for p in pages:
        for ln in p.splitlines():
            s = ln.strip()
            if s in boiler or re.fullmatch(r"(\d{1,3}|[ivxlc]+|SY \d{4}-\d{2}|-----)", s):
                continue
            lines.append(ln.rstrip())
    text = "\n".join(lines)
    text = clean_text(text)
    m = re.search(r"^#+\s*\**\s*(1\.0?|1)\s+\S", text, re.M)
    if not m:
        raise ValueError(f"{pdf.name}: no section 1 heading")
    return text[: m.start()], text[m.start():]


MARK = r"(?:<u>|</u>|\*\*|_)*"


def clean_text(text: str) -> str:
    text = re.sub(rf"\b{MARK}ED{MARK}\s*{MARK}Facts{MARK}", "EDFacts", text)
    text = re.sub(r"</?(u|mark)>", "", text)
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]+$", "", text, flags=re.M)
    text = re.sub(r"^[a-z]$\n?", "", text, flags=re.M)  # stray figure placeholder letters
    text = re.sub(r"^[•·]\s+", "- ", text, flags=re.M)  # bullets pymupdf4llm leaves as text
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text


def relevel(body: str) -> str:
    out = []
    last_numbered = 2
    for ln in body.split("\n"):
        m = HEADING.match(ln)
        if not m:
            out.append(ln)
            continue
        level, title = m.groups()
        title = title.replace("**", "").replace("_", "").strip().rstrip(".").strip()
        title = re.sub(r"\s+", " ", re.sub(r"^[•·]\s*", "", title))
        if not title:
            continue
        n = NUMBERED.match(title)
        if n:
            if n.group(1) is None:
                lvl = 2
            else:
                parts = n.group(1).rstrip(".").split(".")
                if len(parts) == 2 and parts[1] == "0":  # '1.0 Purpose' is a chapter
                    parts = parts[:1]
                lvl = 1 + len(parts)
            last_numbered = lvl
        else:
            lvl = last_numbered + 1
        out.append("#" * min(lvl, 6) + " " + title)
    return "\n".join(out)


def table_after(front: str, marker: str) -> str:
    m = re.search(rf"^#+.*{marker}.*$", front, re.M | re.I)
    if not m:
        return ""
    t = re.search(r"((?:^\|.*\n?)+)", front[m.end():], re.M)
    return t.group(1).strip() if t else ""


def pdf_note(pdf: Path, title: str, fields: dict, linker: Linker, what: str) -> str:
    front, body = pdf_markdown(pdf)
    body = relevel(body)
    body = linker.linkify(body, self_note=fields["title"])
    history = table_after(front, "DOCUMENT HISTORY")
    parts = [frontmatter(fields), "", f"# {title}", "", source_block(fields["source"], what), ""]
    parts.append(
        "> [!note] Screenshots and figures from the PDF are not carried over; headings, text and tables are."
    )
    parts += ["", body.strip()]
    if history:
        parts += ["", "## Document History", "", history]
    return "\n".join(parts) + "\n"


# ----------------------------------------------------------------------------- metadata guides


def metadata_guides(manifest: dict) -> list[dict]:
    """Newest guide per collection from the frozen EDPass metadata page."""
    page = (REF / TARGET_YEAR / "metadata" / "index.html").read_text(encoding="utf-8", errors="ignore")
    local = {Path(p).name: p for p in manifest if p.startswith("metadata/")}
    rows = []
    for sy, sec in re.findall(r"<h3[^>]*>\s*(SY \d{4}-\d{2})(.*?)(?=<h3|</main)", page, re.S):
        for tr in re.findall(r"<tr.*?</tr>", sec, re.S):
            cells = [html.unescape(re.sub(r"<[^>]+>", " ", c)).split() for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.S)]
            hrefs = re.findall(r'href="([^"]+)"', tr)
            if len(cells) < 3 or not hrefs:
                continue
            coll = " ".join(cells[0])
            m = re.match(r"(.*?)\s*v(\d+\.\d+)\.?$", coll)
            name, ver = (m.group(1), "v" + m.group(2)) if m else (coll, "")
            fname = Path(hrefs[0]).name.replace("%20", "_20")
            if fname not in local:
                print(f"!! metadata guide not frozen: {fname}", file=sys.stderr)
                continue
            rows.append({"sy": sy, "collection": name, "version": ver, "fs": [f for f in cells[1] if f.startswith("FS")], "path": local[fname]})
    newest = {}
    for r in sorted(rows, key=lambda r: r["sy"]):
        newest[r["collection"]] = r  # later SY wins
    return sorted(newest.values(), key=lambda r: r["collection"])


# ----------------------------------------------------------------------------- BRSI


def cell(v) -> str:
    s = "" if v is None else str(v).strip()
    if s.lower() in ("none", "nan"):
        return ""
    return s


def md_cell(s: str) -> str:
    s = s.replace("|", "\\|").replace("\r", "")
    s = re.sub(r"\n+", "<br>", s)
    return fix_bold(s)


def brsi_rows(path: Path, sheet: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [cell(h) for h in rows[0]]
    return [dict(zip(hdr, (cell(c) for c in r))) for r in rows[1:] if any(cell(c) for c in r)]


def rule_fs_ids(rule: dict) -> list[str]:
    return ["FS" + x for x in re.findall(r"\b(\d{3})\b", rule["FS"])]


def levels(rule: dict) -> str:
    out = [lv for lv in ("SEA", "LEA", "SCH") if rule.get(lv, "").upper().startswith("Y")]
    if rule.get("Other") and rule["Other"].upper() not in ("N", ""):
        out.append(f"Other ({rule['Other']})")
    return ", ".join(out)


RULE_COLS = ["Rule ID", "Legacy ID", "Phase", "Objective", "Failure", "Levels", "Files", "DGs", "Definition", "Rule logic", "Error message", "Threshold", "Status"]


def rule_row(r: dict, linker: Linker) -> str:
    bare = lambda s, n: re.sub(rf"\b(\d{{{n}}})\b", lambda m: linker.fs("FS" + m.group(1), m.group(1)), s)  # noqa: E731
    files = bare(r["FS"], 3)
    meta = bare(r["Metadata FS"], 4) if r.get("Metadata FS") else ""
    thr = " ".join(x for x in (r.get("Threshold Base", ""), r.get("Threshold %", "")) if x)
    status = r.get("Change Status", "")
    if r.get("Change Explanation"):
        status += f": {r['Change Explanation']}"
    logic = r.get("Rule Logic", "")
    if r.get("Additional Information"):
        logic += f"<br>*{r['Additional Information']}*"
    err = r.get("Error Message", "")
    if r.get("Error Message Condition"):
        err += f"<br>*Condition: {r['Error Message Condition']}*"
    vals = [
        f"**{r['EDPass Rule ID']}**", r.get("Legacy Rule ID", ""), r.get("Data Quality Phase", ""), r.get("Rule Objective", ""),
        r.get("Failure Classification", ""), levels(r), files + (f"<br>Metadata: {meta} DG {r.get('Metadata DG','')}" if meta else ""),
        r.get("DG", ""), linker.linkify(r.get("Definition", "")), linker.linkify(logic), linker.linkify(err), thr, status,
    ]
    return "| " + " | ".join(md_cell(v) for v in vals) + " |"


def rule_table(rules: list[dict], linker: Linker) -> str:
    head = "| " + " | ".join(RULE_COLS) + " |\n|" + "---|" * len(RULE_COLS)
    return head + "\n" + "\n".join(rule_row(r, linker) for r in rules)


def write_brsi(manifest: dict, linker: Linker, specs: dict[str, str]) -> None:
    path = REF / TARGET_YEAR / BRSI_XLSX
    rules = brsi_rows(path, "SY 2025-26")
    discontinued = brsi_rows(path, "Discontinued")
    wb = openpyxl.load_workbook(path, read_only=True)
    sha = manifest.get(BRSI_XLSX, {}).get("sha256", "")
    by_fs = collections.defaultdict(list)
    generic = []
    for r in rules:
        ids = rule_fs_ids(r)
        if ids:
            for i in ids:
                by_fs[i].append(r)
        else:
            generic.append(r)
    disc_by_fs = collections.defaultdict(list)
    for r in discontinued:
        for i in rule_fs_ids(r):
            disc_by_fs[i].append(r)

    # Overview note
    scope = [[cell(c) for c in row] for row in wb["Scope"].iter_rows(values_only=True)][1:]
    hist = [[cell(c) for c in row] for row in wb["Version History"].iter_rows(values_only=True)]
    lines = [
        frontmatter({
            "title": f"EDFacts Business Rules Single Inventory {BRSI_LABEL}", "school_year": "SY 2025-26", "version": "v1.04",
            "rules": len(rules), "discontinued": len(discontinued), "source": BRSI_XLSX,
            "sha256": sha, "frozen": FROZEN, "tags": ["edfacts", "brsi", "business-rules"],
        }),
        "", f"# EDFacts Business Rules Single Inventory — {BRSI_LABEL}", "",
        f"> [!abstract] At a glance\n> **Rules:** {len(rules)} active, {len(discontinued)} discontinued or suspended\n"
        f"> **Per-spec notes:** every converted spec has a `<FSnnn> Business Rules` note listing the rules whose FS column names it "
        f"(e.g. [[FS052 Business Rules]]); the table below links them\n"
        f"> **Not yet published:** a SY {TARGET_YEAR} BRSI. At the {FROZEN} freeze the business-rules page still linked the SY 2025-26 inventory, so it is the working baseline.\n"
        f"> **See also:** [[{INDEX_NOTE}|all specs]], [[{WORKBOOK_NOTE}|workbook]]",
        "", source_block(BRSI_XLSX, "the BRSI workbook"),
        "", "## Rules per file specification", "",
        "| Spec | Active rules | Discontinued | Note |", "|---|---|---|---|",
    ]
    for fs in sorted(set(by_fs) | set(disc_by_fs)):
        note = f"[[{fs} Business Rules]]" if fs in specs else ""
        lines.append(f"| {linker.fs(fs, in_table=True)} | {len(by_fs.get(fs, []))} | {len(disc_by_fs.get(fs, []))} | {note} |")
    lines += ["", "## Scope", "", "The file specifications the inventory covers, per system of publication.", "",
              "| System | Scope | Label |", "|---|---|---|"]
    for row in scope:
        if len(row) >= 2 and any(row):
            lines.append(f"| {row[0]} | {linker.linkify(md_cell(row[1]), in_table=True)} | {row[2] if len(row) > 2 else ''} |")
    lines += ["", "## Rules that apply across files", "",
              "Rules whose FS column names a class of files (\"All files collecting category sets\" ...) rather than specific specs. "
              "They apply to every converted spec in that class.", "", rule_table(generic, linker)]
    lines += ["", "## Version History", "", "| " + " | ".join(hist[0]) + " |", "|" + "---|" * len(hist[0])]
    for row in hist[1:]:
        if any(row):
            lines.append("| " + " | ".join(md_cell(c) for c in row) + " |")
    (OUT / "BRSI").mkdir(parents=True, exist_ok=True)
    (OUT / "BRSI" / f"BRSI SY 2025-26.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    # Per-spec notes
    for fs, note in specs.items():
        active, disc = by_fs.get(fs, []), disc_by_fs.get(fs, [])
        folder = OUT / fs
        folder.mkdir(parents=True, exist_ok=True)
        lines = [
            frontmatter({
                "fs_id": fs, "title": f"{fs} Business Rules", "brsi": BRSI_LABEL, "rules": len(active), "discontinued": len(disc),
                "source": BRSI_XLSX, "sha256": sha, "frozen": FROZEN, "tags": ["edfacts", "brsi", "business-rules"],
            }),
            "", f"# {fs} Business Rules ({BRSI_LABEL})", "",
            f"> [!abstract] At a glance\n> **Spec:** [[{note}]]\n> **Rules naming {fs}:** {len(active)} active, {len(disc)} discontinued or suspended\n"
            f"> **Also applies:** the cross-file rules in [[BRSI SY 2025-26#Rules that apply across files|BRSI SY 2025-26]]\n"
            f"> **Columns:** Phase is the EDPass data-quality phase; Levels is where the rule runs; Files lists every spec the rule reads, so a rule can appear in several notes.",
            "", source_block(BRSI_XLSX, "the BRSI workbook"), "",
            "## Active rules", "", rule_table(active, linker) if active else "_None._",
        ]
        if disc:
            lines += ["", "## Discontinued or suspended rules", "", rule_table(disc, linker)]
        (folder / f"{fs} Business Rules.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"BRSI: {len(rules)} rules, {len(generic)} cross-file, notes for {len(specs)} specs")


# ----------------------------------------------------------------------------- main


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--skip", nargs="*", default=[], choices=["brsi", "workbook", "metadata"])
    args = ap.parse_args()
    manifest = read_manifest(TARGET_YEAR)
    specs = spec_notes()
    guides = metadata_guides(manifest)
    meta_notes = {}
    for g in guides:
        g["note"] = f"{g['collection']} Metadata User Guide"
        for fs in g["fs"]:
            meta_notes.setdefault(fs, g["note"])
    linker = Linker(specs, meta_notes)

    if "metadata" not in args.skip:
        (OUT / "Metadata").mkdir(parents=True, exist_ok=True)
        for g in guides:
            pdf = REF / TARGET_YEAR / g["path"]
            text = pymupdf4llm.to_markdown(str(pdf))
            mentioned = sorted({"FS" + x for x in re.findall(r"\bFS\s?(\d{3})\b", text)} & set(specs))
            fields = {
                "title": g["note"], "collection": g["collection"], "version": g["version"], "school_year": g["sy"],
                "metadata_specs": g["fs"], "related_specs": mentioned, "source": g["path"],
                "sha256": manifest.get(g["path"], {}).get("sha256", ""), "frozen": FROZEN,
                "tags": ["edfacts", "metadata", "edpass"],
            }
            note = pdf_note(pdf, f"{g['collection']} Metadata User Guide ({g['version']}, {g['sy']})", fields, linker, "the EDPass metadata user guide")
            related = ", ".join(linker.fs(f) for f in mentioned) or "none of the converted specs"
            glance = (
                f"> [!abstract] At a glance\n> **Metadata collection:** {g['collection']} {g['version']} ({g['sy']})\n"
                f"> **Metadata file specs:** {', '.join(g['fs'])}\n> **File specs this guide mentions:** {related}\n"
                f"> **See also:** [[{INDEX_NOTE}|all specs]], [[BRSI SY 2025-26|business rules]], [[{WORKBOOK_NOTE}|workbook]]\n"
            )
            note = note.replace("> [!info] Source", glance + "\n> [!info] Source", 1)
            (OUT / "Metadata" / f"{g['note']}.md").write_text(note, encoding="utf-8")
            print(f"metadata: {g['collection']} {g['version']} {g['sy']} -> {len(mentioned)} spec links")

    if "workbook" not in args.skip:
        (OUT / "Workbook").mkdir(parents=True, exist_ok=True)
        pdf = REF / TARGET_YEAR / WORKBOOK_PDF
        fields = {
            "title": WORKBOOK_NOTE, "school_year": "SY 2025-26", "source": WORKBOOK_PDF,
            "sha256": manifest.get(WORKBOOK_PDF, {}).get("sha256", ""), "frozen": FROZEN,
            "tags": ["edfacts", "workbook", "edpass"],
        }
        note = pdf_note(pdf, "EDFacts Workbook with FAQs — SY 2025-26", fields, linker, "the EDFacts Workbook")
        glance = (
            "> [!abstract] At a glance\n> The reference guide for the EDFacts information collection: EDPass roles, the directory "
            "upload order, each program area's files, metadata and data packages, notification templates, FAQs, standard data "
            "definitions, the -1 missing-data rule, race/ethnicity guidance, state codes and the glossary.\n"
            f"> **Newest published:** SY 2025-26 (no SY {TARGET_YEAR} workbook at the {FROZEN} freeze).\n"
            f"> **See also:** [[{INDEX_NOTE}|all specs]], [[BRSI SY 2025-26|business rules]]\n"
        )
        note = note.replace("> [!info] Source", glance + "\n> [!info] Source", 1)
        (OUT / "Workbook" / f"{WORKBOOK_NOTE}.md").write_text(note, encoding="utf-8")
        print("workbook: done")

    if "brsi" not in args.skip:
        write_brsi(manifest, linker, specs)
    return 0


if __name__ == "__main__":
    sys.exit(main())
